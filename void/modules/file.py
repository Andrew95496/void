from .color_code import bcolors as bc

from openpyxl import load_workbook
import pandas as pd


class File():

    __slots__: tuple[str] = ('file_name', 'sheet_name', 'import_dir')

    def __init__(self, file_name: str, sheet_name: str, import_dir: str):
        self.file_name = file_name
        self.sheet_name = sheet_name
        self.import_dir = import_dir

    def import_excel(self) -> list[pd.DataFrame]:
        table_list: list[pd.DataFrame] = []
        workbook = load_workbook(self.import_dir, data_only='True')
        worksheet = workbook[self.sheet_name]

        mapping: dict = {}
        # loop through all the tables and add to a dictionary
        for table, data_boundary in worksheet.tables.items():
            # parse the data within the ref boundary
            data = worksheet[data_boundary]
            ### extract the data ###
            # the inner list comprehension gets the values for each cell in the table
            content = [[cell.value for cell in ent]
                       for ent in data]
            header = content[0]

            # the contents ... excluding the header
            rest = content[1:]

            # create dataframe with the column names
            # and pair table name with dataframe
            df = pd.DataFrame(rest, columns=header)
            mapping[table] = df
            table_list.append(mapping[table])
            self.file_name = self.file_name.replace('_', ' ')
            print(
                f'{bc.HEADER}{bc.BOLD}{self.file_name.capitalize()}{bc.ENDC} ' f'{bc.OKGREEN}{table} Loaded{bc.ENDC}')
            
        return table_list

    def export_excel(self, dataframe: pd.DataFrame, opponent_name: str, export_dir: str) -> None:
        dataframe.to_excel(f'{export_dir}/{opponent_name}.xlsx')
        